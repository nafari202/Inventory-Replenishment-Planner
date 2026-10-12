"""Excel workbook: a formula-driven SKU policy calculator plus the backtest results.

The SKU Policy sheet recalculates ABC class, XYZ class, safety stock, reorder point and order
quantity from the Assumptions sheet, so a planner can change the lead time or a service target
and see every SKU update. check() confirms Excel's numbers match the Python pipeline.
"""
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from . import config

NAVY = "1F3B5C"
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
NOTE_FONT = Font(italic=True, color="595959")
THIN = Side(style="thin", color="D9D9D9")
FIRST = 5  # first data row on table sheets (title, note, blank, header)

GBP = '"£"#,##0'
GBP2 = '"£"#,##0.00'
INT = "#,##0"
DEC1 = "#,##0.0"
PCT1 = "0.0%"


def _title(ws, title, note):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = note
    ws["A2"].font = NOTE_FONT


def _header(ws, row, headers, widths, freeze=True):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 32
    if freeze:  # setting panes and then clearing them leaves XML that Excel has to repair
        ws.freeze_panes = ws.cell(row=row + 1, column=3)


def _name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    _title(ws, "Planning assumptions", "Yellow cells are inputs. Change one and the SKU Policy, ABC-XYZ and Summary sheets recalculate.")
    rows = [
        ("Supplier lead time (weeks)", config.LEAD_TIME_WEEKS, "LeadTime", "0"),
        ("Order cost per purchase order line", config.ORDER_COST, "OrderCost", GBP2),
        ("Annual holding cost (share of unit cost)", config.HOLDING_RATE, "HoldingRate", "0%"),
        ("Unit cost as share of selling price", config.COST_RATIO, "CostRatio", "0%"),
        ("Weeks per year", config.WEEKS_PER_YEAR, "WeeksPerYear", "0"),
        ("Class A cutoff (cumulative revenue share)", config.ABC_CUTOFFS["A"], "CutA", "0%"),
        ("Class B cutoff (cumulative revenue share)", config.ABC_CUTOFFS["B"], "CutB", "0%"),
        ("X cutoff (coefficient of variation at or below)", config.XYZ_CUTOFFS["X"], "CutX", "0.00"),
        ("Y cutoff (coefficient of variation at or below)", config.XYZ_CUTOFFS["Y"], "CutY", "0.00"),
    ]
    _header(ws, 4, ["Assumption", "Value"], [48, 14], freeze=False)
    for i, (label, value, name, fmt) in enumerate(rows, start=FIRST):
        ws.cell(row=i, column=1, value=label)
        c = ws.cell(row=i, column=2, value=value)
        c.fill, c.number_format = INPUT_FILL, fmt
        _name(wb, name, f"Assumptions!$B${i}")
    r = FIRST + len(rows) + 2
    ws.cell(row=r - 1, column=1, value="Cycle service level target by ABC class").font = Font(bold=True, color=NAVY)
    for col, h in enumerate(["Class", "Target"], start=1):
        c = ws.cell(row=r, column=col, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for i, (cls, level) in enumerate(config.SERVICE_TARGETS.items(), start=r + 1):
        ws.cell(row=i, column=1, value=cls)
        c = ws.cell(row=i, column=2, value=level)
        c.fill, c.number_format = INPUT_FILL, "0.0%"
    _name(wb, "TargetClass", f"Assumptions!$A${r + 1}:$A${r + 3}")
    _name(wb, "TargetLevel", f"Assumptions!$B${r + 1}:$B${r + 3}")
    ws.cell(row=r + 5, column=1, value=f"Compared against: safety stock of {config.BASELINE_COVER_WEEKS} weeks of average demand for every SKU, "
            "with the same EOQ order quantities.").font = NOTE_FONT


def sku_policy(wb, prof: pd.DataFrame, planned: pd.Index):
    ws = wb.create_sheet("SKU Policy")
    _title(ws, "SKU replenishment policy", "One row per SKU sold in 2010, ranked by revenue. Grey columns are 2010 data; the rest are formulas.")
    headers = ["Stock code", "Description", "2010 revenue", "Revenue share", "Cumulative share", "ABC", "Avg price",
               "Unit cost", "Avg weekly demand", "Weekly std dev", "CV", "XYZ", "Service target", "z", "Safety stock",
               "Reorder point", "Annual demand", "EOQ", "Order qty", "In 2011 backtest"]
    widths = [11, 34, 12, 10, 11, 6, 9, 9, 11, 11, 8, 6, 10, 7, 10, 10, 11, 9, 9, 10]
    _header(ws, FIRST - 1, headers, widths)
    prof = prof.sort_values("rank")
    last = FIRST + len(prof) - 1
    data_fill = PatternFill("solid", fgColor="F2F2F2")
    for i, (code, p) in enumerate(prof.iterrows(), start=FIRST):
        r = str(i)
        vals = {
            "A": code, "B": p["description"], "C": round(float(p["revenue"]), 2),
            "D": f"=C{r}/SUM($C${FIRST}:$C${last})",
            "E": f"=SUM($C${FIRST}:C{r})/SUM($C${FIRST}:$C${last})",
            "F": f'=IF(E{r}-D{r}<CutA,"A",IF(E{r}-D{r}<CutB,"B","C"))',
            "G": float(p["avg_price"]), "H": f"=G{r}*CostRatio",
            "I": float(p["mean_weekly"]), "J": float(p["sd_weekly"]),
            "K": f"=IF(I{r}>0,J{r}/I{r},0)",
            "L": f'=IF(K{r}<=CutX,"X",IF(K{r}<=CutY,"Y","Z"))',
            "M": f"=_xlfn.XLOOKUP(F{r},TargetClass,TargetLevel)",  # _xlfn. marks post-2007 functions in the file
            "N": f"=_xlfn.NORM.S.INV(M{r})",
            "O": f"=N{r}*J{r}*SQRT(LeadTime)",
            "P": f"=ROUNDUP(ROUND(I{r}*LeadTime+O{r},6),0)",
            "Q": f"=I{r}*WeeksPerYear",
            "R": f"=SQRT(2*Q{r}*OrderCost/(HoldingRate*H{r}))",
            "S": f"=MAX(ROUNDUP(ROUND(R{r},6),0),ROUNDUP(ROUND(I{r},6),0))",
            "T": "Yes" if code in planned else "No",
        }
        for col, v in vals.items():
            ws[f"{col}{r}"] = v
        for col in "ABCGIJ":
            ws[f"{col}{r}"].fill = data_fill
    fmts = {"C": GBP, "D": "0.00%", "E": PCT1, "G": GBP2, "H": GBP2, "I": DEC1, "J": DEC1, "K": "0.00",
            "M": PCT1, "N": "0.00", "O": DEC1, "P": INT, "Q": INT, "R": INT, "S": INT}
    for col, fmt in fmts.items():
        for row in ws[f"{col}{FIRST}:{col}{last}"]:
            row[0].number_format = fmt
    ws.auto_filter.ref = f"A{FIRST - 1}:T{last}"
    return last


def matrix_sheet(wb, last):
    ws = wb.create_sheet("ABC-XYZ")
    _title(ws, "ABC-XYZ matrix", "Counts and revenue shares are COUNTIFS / SUMIFS over the SKU Policy sheet.")
    ws["A4"] = "SKU count"
    ws["A4"].font = Font(bold=True, color=NAVY)
    ws["F4"] = "Share of 2010 revenue"
    ws["F4"].font = Font(bold=True, color=NAVY)
    sp = "'SKU Policy'!"
    rng = lambda col: f"{sp}${col}${FIRST}:${col}${last}"  # noqa: E731
    for block, col0 in (("count", 1), ("share", 6)):
        for j, x in enumerate("XYZ", start=1):
            c = ws.cell(row=5, column=col0 + j, value=x)
            c.fill, c.font = HEAD_FILL, HEAD_FONT
        c = ws.cell(row=5, column=col0 + 4, value="Total")
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        for i, a in enumerate("ABC", start=6):
            ws.cell(row=i, column=col0, value=a).font = Font(bold=True)
            for j, x in enumerate("XYZ", start=1):
                if block == "count":
                    f = f'=COUNTIFS({rng("F")},"{a}",{rng("L")},"{x}")'
                else:
                    f = f'=SUMIFS({rng("C")},{rng("F")},"{a}",{rng("L")},"{x}")/SUM({rng("C")})'
                c = ws.cell(row=i, column=col0 + j, value=f)
                c.number_format = INT if block == "count" else PCT1
            first, lastc = get_column_letter(col0 + 1), get_column_letter(col0 + 3)
            c = ws.cell(row=i, column=col0 + 4, value=f"=SUM({first}{i}:{lastc}{i})")
            c.number_format, c.font = (INT if block == "count" else PCT1), Font(bold=True)
    for col in "ABCDEFGHIJ":
        ws.column_dimensions[col].width = 10


def backtest_sheet(wb, prof, res_pol, res_base, pol, base):
    ws = wb.create_sheet("Backtest")
    _title(ws, "2011 backtest by SKU", "Each policy replayed week by week against actual Jan-Nov 2011 demand. Unmet demand is a lost sale.")
    headers = ["Stock code", "Description", "ABC", "XYZ", "Demand units", "Demand value",
               "ABC policy: reorder point", "ABC policy: order qty", "ABC policy: sold value", "ABC policy: fill rate",
               "ABC policy: stockout weeks", "ABC policy: avg inventory value", "ABC policy: orders",
               "Weeks of cover: reorder point", "Weeks of cover: order qty", "Weeks of cover: sold value", "Weeks of cover: fill rate",
               "Weeks of cover: stockout weeks", "Weeks of cover: avg inventory value", "Weeks of cover: orders"]
    widths = [11, 34, 6, 6, 10, 11, 11, 10, 11, 9, 10, 12, 8, 11, 10, 11, 9, 10, 12, 8]
    _header(ws, FIRST - 1, headers, widths)
    rows = prof.loc[res_pol.index].sort_values("rank")
    last = FIRST + len(rows) - 1
    for i, code in enumerate(rows.index, start=FIRST):
        p, a, b = prof.loc[code], res_pol.loc[code], res_base.loc[code]
        r = str(i)
        vals = [code, p["description"], p["abc"], p["xyz"], a["demand_units"], round(a["demand_value"], 2),
                pol.loc[code, "reorder_point"], pol.loc[code, "order_qty"], round(a["sold_value"], 2),
                f"=IF(F{r}>0,I{r}/F{r},1)", a["stockout_weeks"], round(a["avg_inventory_value"], 2), a["orders"],
                base.loc[code, "reorder_point"], base.loc[code, "order_qty"], round(b["sold_value"], 2),
                f"=IF(F{r}>0,P{r}/F{r},1)", b["stockout_weeks"], round(b["avg_inventory_value"], 2), b["orders"]]
        for col, v in enumerate(vals, start=1):
            ws.cell(row=i, column=col, value=v)
    fmts = {"E": INT, "F": GBP, "G": INT, "H": INT, "I": GBP, "J": PCT1, "K": INT, "L": GBP, "M": INT,
            "N": INT, "O": INT, "P": GBP, "Q": PCT1, "R": INT, "S": GBP, "T": INT}
    for col, fmt in fmts.items():
        for row in ws[f"{col}{FIRST}:{col}{last}"]:
            row[0].number_format = fmt
    red = PatternFill("solid", fgColor="F8CBAD")
    for col in "JQ":
        ws.conditional_formatting.add(f"{col}{FIRST}:{col}{last}", CellIsRule(operator="lessThan", formula=["0.9"], fill=red))
    ws.auto_filter.ref = f"A{FIRST - 1}:T{last}"
    return last


def summary_sheet(wb, last_bt, n_profiled, n_planned):
    ws = wb.create_sheet("Summary", 0)
    _title(ws, "Inventory Replenishment Planner",
           "UCI Online Retail II (UK online retailer). Policy set on 2010 demand, tested on Jan-Nov 2011 demand.")
    bt = "Backtest!"
    rng = lambda col: f"{bt}${col}${FIRST}:${col}${last_bt}"  # noqa: E731
    ws["A4"] = f"{n_profiled:,} SKUs sold in 2010; {n_planned:,} of them also sold in 2011 and are in the backtest."
    headers = ["Policy", "Class", "SKUs", "Demand value", "Sold value", "Fill rate", "Stockout SKU-weeks",
               "Avg inventory value", "Orders placed", "Holding cost", "Ordering cost", "Total cost"]
    widths = [20, 8, 8, 14, 14, 9, 12, 14, 10, 12, 12, 12]
    _header(ws, 6, headers, widths, freeze=False)
    years = len(config.TEST_WEEKS) / config.WEEKS_PER_YEAR
    r = 7
    for label, cols in (("ABC policy", ("I", "K", "L", "M")), (f"Weeks of cover ({config.BASELINE_COVER_WEEKS} wk)", ("P", "R", "S", "T"))):
        sold, so, inv, orders = cols
        for cls in ["All", "A", "B", "C"]:
            crit = "" if cls == "All" else f',{rng("C")},"{cls}"'
            sumif = (lambda col: f"=SUM({rng(col)})") if cls == "All" else (lambda col: f"=SUMIFS({rng(col)}{crit})")
            ws.cell(row=r, column=1, value=label)
            ws.cell(row=r, column=2, value=cls)
            ws.cell(row=r, column=3, value=f"=COUNTA({rng('A')})" if cls == "All" else f'=COUNTIFS({rng("C")},"{cls}")')
            ws.cell(row=r, column=4, value=sumif("F"))
            ws.cell(row=r, column=5, value=sumif(sold))
            ws.cell(row=r, column=6, value=f"=E{r}/D{r}")
            ws.cell(row=r, column=7, value=sumif(so))
            ws.cell(row=r, column=8, value=sumif(inv))
            ws.cell(row=r, column=9, value=sumif(orders))
            ws.cell(row=r, column=10, value=f"=H{r}*HoldingRate*{years:.6f}")
            ws.cell(row=r, column=11, value=f"=I{r}*OrderCost")
            ws.cell(row=r, column=12, value=f"=J{r}+K{r}")
            for col, fmt in zip("CDEFGHIJKL", [INT, GBP, GBP, PCT1, INT, GBP, INT, GBP, GBP, GBP]):
                ws[f"{col}{r}"].number_format = fmt
            if cls == "All":
                for col in range(1, 13):
                    ws.cell(row=r, column=col).font = Font(bold=True)
                    ws.cell(row=r, column=col).border = Border(top=THIN)
            r += 1
    ws.cell(row=r + 1, column=1, value="Fill rate = share of demand value shipped from stock. Holding cost covers the 48 test weeks.").font = NOTE_FONT


def tradeoff_sheet(wb, sweep: pd.DataFrame):
    ws = wb.create_sheet("Trade-off")
    _title(ws, "Service level vs inventory",
           "Each row applies one setting to every SKU and replays 2011 demand. Both methods use the same EOQ order quantities.")
    headers = ["Safety stock method", "Setting", "Fill rate", "Avg inventory value", "Stockout SKU-weeks", "Total cost"]
    _header(ws, 4, headers, [34, 10, 10, 18, 16, 14], freeze=False)
    names = {"Statistical": "Demand variability (service target)", "Weeks of cover": "Weeks of average demand"}
    spans = {}
    r = 5
    for method, g in sweep.groupby("method", sort=False):
        start = r
        for _, s in g.iterrows():
            ws.cell(row=r, column=1, value=names[method])
            ws.cell(row=r, column=2, value=s["label"])
            for col, (key, fmt) in enumerate([("fill_rate", PCT1), ("avg_inventory_value", GBP),
                                              ("stockout_weeks", INT), ("total_cost", GBP)], start=3):
                c = ws.cell(row=r, column=col, value=float(s[key]))
                c.number_format = fmt
            r += 1
        spans[names[method]] = (start, r - 1)
    chart = ScatterChart()
    chart.title = "Fill rate vs average inventory value"
    chart.scatterStyle = "lineMarker"
    chart.title.overlay = False
    chart.y_axis.number_format = "0%"
    chart.x_axis.number_format = '"£"#,##0,"K"'
    chart.x_axis.delete = chart.y_axis.delete = False  # openpyxl hides axis labels unless told otherwise
    fills = sweep["fill_rate"].astype(float)
    chart.y_axis.scaling.min = float(np.floor(fills.min() * 50) / 50)
    chart.y_axis.scaling.max = float(np.ceil(fills.max() * 50) / 50)
    chart.y_axis.majorUnit = 0.01
    chart.x_axis.scaling.min = float(np.floor(sweep["avg_inventory_value"].min() / 100000) * 100000)
    chart.legend.position = "b"
    chart.legend.overlay = False
    for axis in (chart.x_axis, chart.y_axis):
        axis.majorGridlines = axis.majorGridlines or ChartLines()
        axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9"))
    colors = {"Demand variability (service target)": "2A78D6", "Weeks of average demand": "898781"}
    for title, (a, b) in spans.items():
        series = Series(Reference(ws, min_col=3, min_row=a, max_row=b), Reference(ws, min_col=4, min_row=a, max_row=b), title=title)
        series.marker.symbol = "circle"
        series.marker.size = 6
        series.marker.graphicalProperties.solidFill = colors[title]
        series.marker.graphicalProperties.line.solidFill = colors[title]
        series.graphicalProperties.line.solidFill = colors[title]
        series.graphicalProperties.line.width = 28575  # 2.25 pt
        series.smooth = False
        chart.series.append(series)
    chart.height, chart.width = 10, 18
    ws.add_chart(chart, "H4")


def build(path, prof, planned, res_pol, res_base, pol, base, sweep):
    wb = Workbook()
    wb.remove(wb.active)
    assumptions(wb)
    last_sku = sku_policy(wb, prof, planned)
    matrix_sheet(wb, last_sku)
    last_bt = backtest_sheet(wb, prof, res_pol, res_base, pol, base)
    tradeoff_sheet(wb, sweep)
    summary_sheet(wb, last_bt, len(prof), len(planned))
    wb.active = 0
    wb.save(path)
    return path


def check(path, prof: pd.DataFrame, pol_all: pd.DataFrame) -> dict:
    """Compare Excel's calculated values (after an Excel recalc) with the Python pipeline."""
    ws = load_workbook(path, data_only=True)["SKU Policy"]
    got = {}
    for row in ws.iter_rows(min_row=FIRST, values_only=True):
        if row[0] is None:
            break
        got[row[0]] = {"abc": row[5], "xyz": row[11], "ss": row[14], "rop": row[15], "q": row[18]}
    got = pd.DataFrame(got).T
    if got["abc"].isna().all():
        return {"calculated": False}
    p = prof.loc[got.index]
    q = pol_all.loc[got.index]
    return {
        "calculated": True,
        "skus": len(got),
        "abc_match": int((got["abc"] == p["abc"]).sum()),
        "xyz_match": int((got["xyz"] == p["xyz"]).sum()),
        "safety_stock_match": int(((got["ss"].astype(float) - q["safety_stock"]).abs() < 1e-6).sum()),
        "reorder_point_match": int((got["rop"].astype(float) == q["reorder_point"]).sum()),
        "order_qty_match": int((got["q"].astype(float) == q["order_qty"]).sum()),
    }
